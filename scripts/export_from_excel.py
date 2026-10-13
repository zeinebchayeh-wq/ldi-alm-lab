"""
Day 1: export the model inputs from the SOA Excel tool to CSV files.

Why: the Excel file is our "answer key". We copy its INPUTS once into simple
CSV files, so our Python model never depends on Excel again. Later we will
compare our OUTPUTS with the Excel outputs.

Usage:
    python scripts/export_from_excel.py path/to/liability-driven-investment.xlsm
"""
import sys
from pathlib import Path

import openpyxl
import pandas as pd

DATA_DIR = Path(__file__).resolve().parents[1] / "data"


def read_block(ws, first_row, last_row, first_col, last_col):
    """Read a rectangular block of cells and return a list of rows (lists).

    Rows and columns are 1-based, like in Excel (column A = 1, B = 2, ...).
    """
    rows = []
    for row in ws.iter_rows(min_row=first_row, max_row=last_row,
                            min_col=first_col, max_col=last_col,
                            values_only=True):
        rows.append(list(row))
    return rows


def export_census(wb):
    """Census sheet: one row per plan participant (rows 4-13, columns A-I)."""
    ws = wb["Census"]
    header = read_block(ws, 3, 3, 1, 9)[0]          # row 3 = column titles
    body = read_block(ws, 4, 13, 1, 9)               # rows 4-13 = 10 participants
    census = pd.DataFrame(body, columns=header)
    # Keep only the date part (Excel stores dates as datetimes at midnight)
    for col in ["Date of Birth", "Date of Hire", "Date of Retirement"]:
        census[col] = pd.to_datetime(census[col]).dt.date
    # Short, code-friendly column names
    census.columns = ["id", "retired", "gender", "birth_date", "hire_date",
                      "salary", "retirement_date", "occupation", "weight"]
    return census


def export_assumptions(wb):
    """Actuarial Assumption sheet: several small tables side by side."""
    ws = wb["Actuarial Assumption"]

    # 1) Base mortality table q_x: ages 0-110 in rows 5-115, columns A-C
    mortality = pd.DataFrame(read_block(ws, 5, 115, 1, 3),
                             columns=["age", "qx_male", "qx_female"])

    # 2) Turnover rate during the vesting period, by occupation (E5:F9)
    turnover = pd.DataFrame(read_block(ws, 5, 9, 5, 6),
                            columns=["occupation", "turnover_rate"])

    # 3) Salary growth multiple by occupation (E13:F17)
    salary_multiple = pd.DataFrame(read_block(ws, 13, 17, 5, 6),
                                   columns=["occupation", "salary_growth_multiple"])

    # 4) Scalar assumptions, stored as one "name, value" table
    def cell(ref):
        return ws[ref].value

    scalars = pd.DataFrame([
        ("mortality_improvement_male", cell("E21")),
        ("mortality_improvement_female", cell("F21")),
        ("shock_mortality", cell("I4")),
        ("shock_mortality_improvement", cell("I5")),
        ("shock_turnover", cell("I6")),
        ("shock_salary_growth", cell("I7")),
        ("share_lump_sum", cell("I12")),
        ("share_life_payment", cell("I13")),
    ], columns=["name", "value"])

    return mortality, turnover, salary_multiple, scalars


def main(xlsm_path):
    # data_only=True -> we read the VALUES Excel last calculated, not formulas
    wb = openpyxl.load_workbook(xlsm_path, read_only=True, data_only=True)
    DATA_DIR.mkdir(exist_ok=True)

    census = export_census(wb)
    census.to_csv(DATA_DIR / "census.csv", index=False)

    mortality, turnover, salary_multiple, scalars = export_assumptions(wb)
    mortality.to_csv(DATA_DIR / "mortality.csv", index=False)
    turnover.to_csv(DATA_DIR / "turnover.csv", index=False)
    salary_multiple.to_csv(DATA_DIR / "salary_multiple.csv", index=False)
    scalars.to_csv(DATA_DIR / "assumptions.csv", index=False)

    print(f"Exported {len(census)} participants and {len(mortality)} mortality ages to {DATA_DIR}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit("Usage: python scripts/export_from_excel.py path/to/file.xlsm")
    main(sys.argv[1])
